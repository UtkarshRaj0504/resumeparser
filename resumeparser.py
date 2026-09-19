import os 
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
import time

load_dotenv()
my_api_key=os.getenv("groq_apikey")
if not my_api_key:
    raise ValueError( "api key not found")

client=Groq(api_key=my_api_key)
model ="openai/gpt-oss-120b"

job_description="""Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.

Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.

Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older

Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills

Amazon is an equal opportunity employer and does not discriminate on the basis of protected veteran status, disability, or other legally protected status.

Our inclusive culture empowers Amazonians to deliver the best results for our customers. If you have a disability and need a workplace accommodation or adjustment during the application and hiring process, including support for the interview or onboarding process, please visit https://amazon.jobs/content/en/how-we-hire/accommodations for more information. If the country/region you’re applying in isn’t listed, please contact your Recruiting Partner.


The base salary range for this position is listed below. Your Amazon package will include sign-on payments and restricted stock units (RSUs). Final compensation will be determined based on factors including experience, qualifications, and location. Amazon also offers comprehensive benefits including health insurance (medical, dental, vision, prescription, Basic Life & AD&D insurance and option for Supplemental life plans, EAP, Mental Health Support, Medical Advice Line, Flexible Spending Accounts, Adoption and Surrogacy Reimbursement coverage), 401(k) matching, paid time off, and parental leave. Learn more about our benefits at https://amazon.jobs/en/benefits.
"""

class jobd(BaseModel):
    role:str
    required_skills:str
    preferred_skills:list[str]
    minimum_experience:float | None
    education_requirements: list[str]
    responsibilities: list[str]

jobd_schema = jobd.model_json_schema()


def get_job(job_description_text=None):
    system_prompt = f"""
    You are an expert HR assistant.

    Your job is to analyse a job description and extract structured
    information from it.

    Return ONLY valid JSON matching this schema:

    {jobd_schema}

    IMPORTANT:
    - Do not return the schema itself.
    - Fill the fields with actual information from the job description.
    - If minimum experience is not mentioned, return null.
    - If information for a list is missing, return an empty list.
    - Do not invent information.
    """

    user_prompt = f"""
    Analyse the following job description:

    {job_description_text if job_description_text else job_description}
    """

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        temperature=2,
        response_format={"type": "json_object"}
    )

    raw_json = response.choices[0].message.content
    job_data = json.loads(raw_json)

    return jobd(**job_data)
response_format={
    "type":"json_object"}
system_prompt=f"""you are an expert hr assistant. your job is to analyse job description and extract structured information from them return only valid JSON matching this schema {jobd_schema}
IMPORTANT: do not return the schema itself. do not return field like "properties","title" or "type" .fill the schema with actual information extracted from the job description.
if minimum experience is not mentioned then return null.
if information for a list is missing then return empty list.
do not invent information."""
message_system=[{
    "role":"system",
    "content":system_prompt}]

user_prompt=f"""analyze the following job description:
{job_description}"""
message=[{
    "role":"user",
    "content":user_prompt}]

messages=message_system+message

response=client.chat.completions.create(model=model,messages=messages,temperature=2,response_format=response_format)
print("#########################################")
answer=response.choices[0].message.content
print(answer)

import json
raw_json=answer
job_data=json.loads(raw_json)
job=jobd(**job_data)
print(job.minimum_experience)
print(job.education_requirements)

class MatchResult(BaseModel):
    score: float
    details: dict
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()
def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



